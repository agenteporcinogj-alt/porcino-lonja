#!/usr/bin/env python3
# PIPELINE NUBE - Grupo Jorge / precio cerdo Lonja de Lleida
# correo -> PDFs -> Excel maestro -> modelo -> prediccion -> registro -> web cifrada
import os, re, csv, json, base64, datetime, statistics as st, glob, imaplib, email
from email.header import decode_header
import pdfplumber, openpyxl
from openpyxl.styles import PatternFill
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

HERE=os.path.dirname(os.path.abspath(__file__))
PDFS=os.path.join(HERE,'_pdfs')
MASTER=os.path.join(HERE,'data','master.xlsx')
LOG=os.path.join(HERE,'data','validacion_modelo.csv')
ENC_OUT=os.path.join(HERE,'docs','data.enc.js')
NUM=re.compile(r'-?\d+,\d+')
def tf(s): return float(s.replace('.','').replace(',','.'))
def nbe(l): return [tf(x) for x in NUM.findall(re.split(r'€',l)[0])]

# ---------- 1) CORREO (ultimos ~30 dias) ----------
def paso_correo():
    user=os.environ.get('GMAIL_USER'); pw=os.environ.get('GMAIL_APP_PASSWORD')
    os.makedirs(PDFS,exist_ok=True)
    if not user or not pw:
        print('  (sin credenciales de correo, uso lo que haya en disco)'); return
    since=(datetime.date.today()-datetime.timedelta(days=30)).strftime('%d-%b-%Y')
    M=imaplib.IMAP4_SSL('imap.gmail.com'); M.login(user,pw); M.select('INBOX')
    typ,data=M.search(None,f'(SINCE {since})'); ids=data[0].split()
    saved=0
    for i in ids:
        typ,md=M.fetch(i,'(RFC822)'); msg=email.message_from_bytes(md[0][1])
        for part in msg.walk():
            if part.get_content_maintype()=='multipart': continue
            fn=part.get_filename()
            if not fn or not fn.lower().endswith('.pdf'): continue
            parts=decode_header(fn); nm=''
            for txt,enc in parts: nm+= txt.decode(enc or 'utf-8','ignore') if isinstance(txt,bytes) else txt
            nm=re.sub(r'[^A-Za-z0-9._\- ]','_',nm).strip() or 'adj.pdf'
            path=os.path.join(PDFS,nm)
            if os.path.exists(path): continue
            with open(path,'wb') as f: f.write(part.get_payload(decode=True)); saved+=1
    M.logout(); print(f'  PDFs nuevos del correo: {saved}')

# ---------- extraccion ----------
def isoweek(fn,pfx):
    m=re.match(pfx+r'(\d{2})(\d{2})(\d{2})',fn)
    if not m: return None
    try: d=datetime.date(2000+int(m.group(1)),int(m.group(2)),int(m.group(3)))
    except: return None
    return d.isocalendar()[:2]
def index_files(pfx):
    out={}
    for f in glob.glob(os.path.join(PDFS,pfx+'[0-9]*.pdf')):
        yw=isoweek(os.path.basename(f),pfx)
        if yw and yw[0]==2026 and yw[1] not in out: out[yw[1]]=f
    return out
def esp_de_po(path):
    # Objetivo = precio REAL de la Lonja de Lleida ("Cerdo Blanco", valor de esta semana).
    try:
        with pdfplumber.open(path) as pdf: allt='\n'.join((p.extract_text() or '') for p in pdf.pages)
    except: return None
    for l in allt.split('\n'):
        s=l.strip().lower()
        if s.startswith('cerdo blanco') or s.startswith('cerdo cebado'):
            ns=NUM.findall(l)
            if len(ns)>=2: return tf(ns[1])
    for l in allt.split('\n'):  # fallback: equivalente España
        if l.strip().startswith('España'):
            ns=NUM.findall(l)
            if len(ns)>=12: return tf(ns[6])
    return None
def ale_fra_de_de(path):
    try:
        with pdfplumber.open(path) as pdf: t=pdf.pages[0].extract_text() or ''
    except: return None,None
    ale=fra=None
    for l in t.split('\n'):
        u=l.upper()
        if 'NW-AMI' in u and 'CANAL' in u:
            nb=nbe(l); ale=nb[-1] if nb else ale
        if l.strip().upper().startswith('FRANCIA') and ('MPB' in u or 'MPF' in u):
            nb=nbe(l); fra=nb[-1] if nb else fra
    return ale,fra
def din_de_ip(path):
    try:
        with pdfplumber.open(path) as pdf: t=pdf.pages[0].extract_text() or ''
    except: return None
    if 'EUROS KILO VIVO' not in t.upper(): return None
    for l in t.split('\n'):
        if l.strip().startswith('Dinamarca'):
            ns=NUM.findall(l)
            if len(ns)>=2: return tf(ns[0])
    return None

# ---------- 2) ACTUALIZAR MAESTRO ----------
def paso_maestro():
    po=index_files('PO'); de=index_files('DE'); ip=index_files('IP')
    wb=openpyxl.load_workbook(MASTER)
    YEL=PatternFill('solid',fgColor='FFF2CC'); ORA=PatternFill('solid',fgColor='FCE4D6')
    def g(sh,w,col): return wb[sh].cell(8+w,col).value
    nuevas=[]
    for w in range(1,53):
        r=8+w
        if g('CEBADO LLEIDA',w,18) is None and w in po:
            eq=esp_de_po(po[w])
            if eq is not None: wb['CEBADO LLEIDA'].cell(r,18).value=eq; wb['CEBADO LLEIDA'].cell(r,18).fill=YEL; nuevas.append(w)
        if w in de:
            ale,fra=ale_fra_de_de(de[w])
            if g('ALEMANIA',w,18) is None and ale is not None: wb['ALEMANIA'].cell(r,18).value=ale; wb['ALEMANIA'].cell(r,18).fill=YEL
            if g('FRANCIA',w,35) is None and fra is not None: wb['FRANCIA'].cell(r,35).value=fra; wb['FRANCIA'].cell(r,35).fill=YEL
        if g('DINAMARCA',w,18) is None and w in ip:
            d=din_de_ip(ip[w])
            if d is not None: wb['DINAMARCA'].cell(r,18).value=d; wb['DINAMARCA'].cell(r,18).fill=YEL
    for sheet,col in [('ALEMANIA',18),('FRANCIA',35)]:
        lw=max([w for w in range(1,53) if g(sheet,w,col) is not None] or [0])
        for w in range(2,lw+1):
            c=wb[sheet].cell(8+w,col)
            if c.value is None and wb[sheet].cell(8+w-1,col).value is not None:
                c.value=wb[sheet].cell(8+w-1,col).value; c.fill=ORA
    wb.save(MASTER)
    print(f'  Semanas España nuevas en maestro: {nuevas}')

# ---------- 3) MODELO ----------
def solve(A,b):
    n=len(A); M=[row[:]+[b[i]] for i,row in enumerate(A)]
    for c in range(n):
        p=max(range(c,n),key=lambda r:abs(M[r][c])); M[c],M[p]=M[p],M[c]; piv=M[c][c]
        if abs(piv)<1e-12: continue
        for r in range(n):
            if r!=c and abs(M[r][c])>1e-12:
                f=M[r][c]/piv
                for k in range(c,n+1): M[r][k]-=f*M[c][k]
    return [M[i][n]/M[i][i] if abs(M[i][i])>1e-12 else 0.0 for i in range(n)]
def fit(data):
    k=len(data[0]['x']); ATA=[[0.0]*k for _ in range(k)]; ATy=[0.0]*k
    for d in data:
        for i in range(k):
            ATy[i]+=d['x'][i]*d['dESP']
            for j in range(k): ATA[i][j]+=d['x'][i]*d['x'][j]
    return solve(ATA,ATy)
def paso_modelo():
    wb=openpyxl.load_workbook(MASTER,data_only=True); years=list(range(2010,2027))
    def cell(sh,w,col):
        v=wb[sh].cell(8+w,col).value; return v if isinstance(v,(int,float)) else None
    seq=[]
    for y in years:
        yi=years.index(y)
        for w in range(1,53):
            e=cell('CEBADO LLEIDA',w,2+yi)
            if e is None: continue
            seq.append({'y':y,'w':w,'e':e,'f':cell('FRANCIA',w,2+2*yi+1),'a':cell('ALEMANIA',w,2+yi),'d':cell('DINAMARCA',w,2+yi)})
    esp={(r['y'],r['w']):r['e'] for r in seq}; allmean=st.mean(esp.values())
    seas={w: st.mean([esp[(y,w)] for y in years if (y,w) in esp])/allmean*100 for w in range(1,53)}
    rows=[]
    for t in range(2,len(seq)):
        c,p1,p2=seq[t],seq[t-1],seq[t-2]
        rows.append({'y':c['y'],'w':c['w'],'prev':p1['e'],'dESP':c['e']-p1['e'],'target':c['e'],
            'x':[1.0,p1['e']-p2['e'],(p1['f']-p2['f']) if p1['f'] and p2['f'] else 0.0,
                 (p1['a']-p2['a']) if p1['a'] and p2['a'] else 0.0,(seas[c['w']]-seas[p1['w']])/100.0*allmean]})
    em=[];en=[];within=0;nb=0
    for d in rows:
        if d['y']<2019: continue
        tr=[r for r in rows if (r['y']*100+r['w'])<(d['y']*100+d['w'])]
        if len(tr)<60: continue
        b=fit(tr); pr=d['prev']+sum(bi*xi for bi,xi in zip(b,d['x']))
        em.append(abs(pr-d['target'])); en.append(abs(d['prev']-d['target']))
        if abs(pr-d['target'])<=0.02: within+=1
        nb+=1
    beta=fit(rows); last=seq[-1]; p1=seq[-2]; p2=seq[-3]; nextw=last['w']%52+1
    x=[1.0,last['e']-p1['e'],(last['f']-p1['f']) if last['f'] and p1['f'] else 0.0,
       (last['a']-p1['a']) if last['a'] and p1['a'] else 0.0,(seas[nextw]-seas[last['w']])/100.0*allmean]
    pred=round(last['e']+sum(bi*xi for bi,xi in zip(beta,x)),3)
    contrib={'inercia':round(beta[1]*x[1]*100,1),'francia':round(beta[2]*x[2]*100,1),
             'alemania':round(beta[3]*x[3]*100,1),'estacional':round(beta[4]*x[4]*100,1)}
    vecinos={'francia_delta':round((last['f']-p1['f'])*100,1) if last['f'] and p1['f'] else None,
             'alemania_delta':round((last['a']-p1['a'])*100,1) if last['a'] and p1['a'] else None}
    serie=[{'y':r['y'],'w':r['w'],'v':r['e']} for r in seq[-104:]]
    win=seq[-104:]
    def idx(win,key):
        base=next((r[key] for r in win if r.get(key) is not None),None)
        return [round(r[key]/base*100,1) if (r.get(key) is not None and base) else None for r in win]
    serie_paises={'labels':[{'y':r['y'],'w':r['w']} for r in win],
                  'esp':idx(win,'e'),'fra':idx(win,'f'),'ale':idx(win,'a'),'din':idx(win,'d')}
    return {'mae_m':sum(em)/len(em),'mae_n':sum(en)/len(en),'within2':round(within/nb*100) if nb else 0,
            'last':{'y':last['y'],'w':last['w'],'v':round(last['e'],3)},
            'nextw':nextw,'nexty':last['y'] if nextw>last['w'] else last['y']+1,'pred':pred,'esp':esp,'serie':serie,
            'delta_cts':round((pred-last['e'])*100,1),'contrib':contrib,'vecinos':vecinos,
            'seas':{str(w):round(seas[w],1) for w in seas},'serie_paises':serie_paises}

# ---------- 4) REGISTRO ----------
def paso_registro(m):
    L={}
    if os.path.exists(LOG):
        for r in csv.DictReader(open(LOG)): L[(int(r['anio']),int(r['semana']))]=r
    for (y,w),r in L.items():
        if (not r.get('real')) and (y,w) in m['esp']:
            real=m['esp'][(y,w)]; r['real']=f"{real:.3f}"; r['error_cts']=f"{abs(float(r['prediccion'])-real)*100:.1f}"
    key=(m['nexty'],m['nextw'])
    if key not in L:
        L[key]={'anio':m['nexty'],'semana':m['nextw'],'fecha_pred':datetime.date.today().isoformat(),
                'prediccion':f"{m['pred']:.3f}",'real':'','error_cts':''}
    with open(LOG,'w',newline='') as fh:
        w=csv.DictWriter(fh,fieldnames=['anio','semana','fecha_pred','prediccion','real','error_cts']); w.writeheader()
        for k in sorted(L): w.writerow(L[k])
    return [L[k] for k in sorted(L)]

# ---------- 5) CIFRAR PARA LA WEB ----------
def cifrar(plaintext,password):
    salt=os.urandom(16)
    kdf=PBKDF2HMAC(algorithm=hashes.SHA256(),length=32,salt=salt,iterations=200000)
    key=kdf.derive(password.encode())
    iv=os.urandom(12); ct=AESGCM(key).encrypt(iv,plaintext.encode(),None)
    return base64.b64encode(salt+iv+ct).decode()
def paso_web(m,historial):
    delta=m['delta_cts']; base=m['pred']
    if delta>=1.0: sem={'estado':'AGUANTA','color':'verde','txt':'El precio va a SUBIR ~%.1f cts la semana que viene. Si puedes, aguanta la venta.'%delta}
    elif delta<=-1.0: sem={'estado':'VENDE','color':'rojo','txt':'El precio va a BAJAR ~%.1f cts la semana que viene. Vender ahora protege margen.'%abs(delta)}
    else: sem={'estado':'ESTABLE','color':'ambar','txt':'Precio estable (±%.1f cts). Sin presión para adelantar ni retrasar ventas.'%abs(delta)}
    escenarios=[
      {'nombre':'Base (modelo)','valor':round(base,2),'txt':'Lo más probable con los datos actuales.'},
      {'nombre':'Se agrava PPA / cierran mercados','valor':round(base-0.15,2),'txt':'Menos exportación, sobra carne → precio abajo. (~-15 cts, ilustrativo)'},
      {'nombre':'China reabre / sube demanda','valor':round(base+0.15,2),'txt':'Más demanda exterior → precio arriba. (~+15 cts, ilustrativo)'}]
    eventos=[{'fecha':'6-7 oct 2026','titulo':'Juicio de densidades (Aragón)','txt':'Un juez decide si se aplica la norma europea de densidades (cortar colas → más espacio por animal). Si sale adelante: menos cerdos por granja → menos oferta española (puede empujar el precio arriba) y ~-10% de capacidad para Grupo Jorge (~740 cebos).'}]
    payload={'generado':datetime.datetime.now().strftime('%d/%m/%Y %H:%M'),
             'ultimo':m['last'],'prediccion':{'y':m['nexty'],'w':m['nextw'],'v':m['pred']},'delta_cts':delta,
             'precision':{'modelo_cts':round(m['mae_m']*100,1),'naive_cts':round(m['mae_n']*100,1),'within2':m['within2']},
             'historial':[{'semana':r['semana'],'pred':r['prediccion'],'real':r.get('real',''),'error':r.get('error_cts','')} for r in historial],
             'serie':m['serie'],'serie_paises':m['serie_paises'],'contrib':m['contrib'],'vecinos':m['vecinos'],
             'seas':m['seas'],'semaforo':sem,'escenarios':escenarios,'eventos':eventos}
    pw=os.environ.get('WEB_PASSWORD','cerdo')
    enc=cifrar(json.dumps(payload,ensure_ascii=False),pw)
    open(ENC_OUT,'w').write('window.ENC="'+enc+'";')
    print(f'  Web cifrada actualizada ({len(enc)} bytes).')

if __name__=='__main__':
    print('== 1) Correo =='); paso_correo()
    print('== 2) Maestro =='); paso_maestro()
    print('== 3) Modelo =='); m=paso_modelo()
    print('== 4) Registro =='); h=paso_registro(m)
    print('== 5) Web =='); paso_web(m,h)
    print(f"\nOK. Último {m['last']['w']}={m['last']['v']} | Predicción sem {m['nextw']}={m['pred']} | error {round(m['mae_m']*100,1)} cts")
